"""
Convert docs/MANUSCRIPT.md into a formatted Word document (docs/MANUSCRIPT.docx).
Handles: H1/H2/H3 headings, horizontal rules, bullet lists, numbered lists,
GitHub-style tables, inline **bold** and *italic*, and normal paragraphs.
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

HERE = Path(__file__).resolve().parent
SRC = HERE / "MANUSCRIPT.md"
OUT = HERE / "MANUSCRIPT.docx"

INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")


def add_runs(paragraph, text):
    """Add text to a paragraph, honoring **bold** and *italic* markers."""
    # Protect escaped asterisks (\*) so they are not parsed as emphasis.
    text = text.replace(r"\*", "\x00")
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2].replace("\x00", "*"))
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1].replace("\x00", "*"))
            run.italic = True
        else:
            paragraph.add_run(part.replace("\x00", "*"))


def is_table_row(line):
    return line.strip().startswith("|") and line.strip().endswith("|")


def parse_table_row(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()

    # Base style
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    i = 0
    n = len(lines)
    while i < n:
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        # Blank line
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if stripped == "---":
            # Represent as a thin empty spacer paragraph
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # Headings
        if stripped.startswith("# "):
            h = doc.add_heading(level=0)
            add_runs(h, stripped[2:].strip())
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue

        # Tables (consecutive | ... | lines)
        if is_table_row(line):
            table_lines = []
            while i < n and is_table_row(lines[i]):
                table_lines.append(lines[i])
                i += 1
            # Drop separator rows like |---|---|
            rows = [parse_table_row(l) for l in table_lines
                    if not re.match(r"^\s*\|[\s:|-]+\|\s*$", l)]
            if rows:
                ncols = max(len(r) for r in rows)
                table = doc.add_table(rows=0, cols=ncols)
                table.style = "Light Grid Accent 1"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ridx, row in enumerate(rows):
                    cells = table.add_row().cells
                    for cidx in range(ncols):
                        text = row[cidx] if cidx < len(row) else ""
                        cell_p = cells[cidx].paragraphs[0]
                        add_runs(cell_p, text)
                        if ridx == 0:  # header row bold
                            for r in cell_p.runs:
                                r.bold = True
            continue

        # Bullet list
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, stripped[2:].strip())
            i += 1
            continue

        # Numbered list
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, m.group(2).strip())
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_runs(p, stripped)
        i += 1

    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
